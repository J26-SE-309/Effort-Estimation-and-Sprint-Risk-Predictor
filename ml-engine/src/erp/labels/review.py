"""The 200-story human check of the automatic labels (ML guide 6.3, objective SO1).

Draws a fixed sample of labelled stories and writes one workbook per reviewer. Each row shows what
happened to the story around its first sprint (sprint moves, status, resolution, story points, blocking
links) with a link to the public Jira page, and leaves yellow cells for the reviewer's verdict. The
automatic labels are kept out of the workbooks, in a separate key file, so reviewers are not steered
by them. Reviewers must be people: the agreement reported in the thesis is between humans.

Usage:
    erp-review-sample                       # writes ERP_DATA_DIR/effort-risk/review/
"""

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
import pyarrow.compute as pc
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from erp import config, tawos
from erp.labels import rules

RARE_RULES = ("r2", "r3", "r4", "r6")
REVIEWERS = ("A", "B")
MAX_EVENTS = 40
FONT = "Arial"

SIGNS = [
    ("R1", "Spillover", "It was in the sprint at the end but not finished (not resolved) when the sprint closed."),
    ("R2", "Delayed closure", "It was finished only after the sprint ended, while still counted in that sprint."),
    ("R3", "Re-estimation", "Its story points were changed while the sprint was running."),
    ("R4", "Blocked", "It waited on another unfinished issue ('is blocked by' / 'depends on') for more than "
                      "about 30% of the sprint."),
    ("R5", "Carry-over", "It was moved into the next sprint unfinished."),
    ("R6", "Reopening", "It was marked done, then reopened in the same sprint or the next one."),
]


def sample_for_review(labels: pd.DataFrame, n: int = 200, per_rule: int = 10, seed: int = 42) -> pd.DataFrame:
    """Half at-risk, half not-at-risk stories by the automatic labels, with the rare rules guaranteed.

    The at-risk half first takes `per_rule` stories for each rarely firing rule (R2, R3, R4, R6) so the
    review can check them at all, then fills up at random; the other half is random. The order is shuffled
    so the two groups are not in blocks. The same seed always gives the same sample.
    """
    half = n // 2
    risky, safe = labels[labels["at_risk"]], labels[~labels["at_risk"]]
    picked = []
    for i, rule in enumerate(RARE_RULES):
        pool = risky[risky[rule] & ~risky["Issue_ID"].isin([p for chunk in picked for p in chunk])]
        picked.append(pool.sample(min(per_rule, len(pool)), random_state=seed + i)["Issue_ID"].tolist())
    chosen = [p for chunk in picked for p in chunk]
    rest = risky[~risky["Issue_ID"].isin(chosen)].sample(half - len(chosen), random_state=seed)
    at_risk_half = risky[risky["Issue_ID"].isin(chosen + rest["Issue_ID"].tolist())].copy()
    at_risk_half["stratum"] = np.where(at_risk_half["Issue_ID"].isin(chosen), "at risk (rare rule)", "at risk")
    sample = pd.concat([at_risk_half, safe.sample(n - half, random_state=seed).assign(stratum="not at risk")])
    return sample.sample(frac=1, random_state=seed).reset_index(drop=True)


def browse_url(api_url: str | None, key: str) -> str | None:
    """TAWOS stores REST API URLs (https://host/jira/rest/api/2/issue/123); people need /browse/KEY."""
    if not isinstance(api_url, str) or "/rest/api/" not in api_url:
        return None
    return api_url.split("/rest/api/")[0] + "/browse/" + key


def describe(field: str, old, new) -> str | None:
    """One line of the story's history in plain words, or None for changes the reviewer does not need."""
    old = old if isinstance(old, str) and old.strip() else None
    new = new if isinstance(new, str) and new.strip() else None
    if old == new:
        return None
    if field == "Sprint":
        return f"Sprint field: {old or 'none'} → {new or 'none'}"
    if field == "status":
        return f"Status: {old or '?'} → {new or '?'}"
    if field == "resolution":
        return f"Resolved as {new}" if new else f"Resolution cleared (was {old})"
    if field == "Story Points":
        return f"Story points: {old or 'none'} → {new or 'none'}"
    if field == "Link":
        for text, verb in ((new, "Link added"), (old, "Link removed")):
            if text and rules.BLOCKING_LINK.match(text):
                return f"{verb}: {text}"
    return None


def history_text(events: pd.DataFrame, story, limit: int = MAX_EVENTS) -> str:
    """The story's events from three days before its first sprint to the end of the following sprint."""
    length = story.sprint_closed - story.sprint_start
    start = min(story.sprint_start, story.commitment_time) - pd.Timedelta(days=3)
    end = story.sprint_closed + length
    fmt = "%Y-%m-%d %H:%M"
    lines = [(story.sprint_start, f"▶ First sprint starts: {story.Sprint_Name}"),
             (story.sprint_closed, f"■ First sprint closed: {story.Sprint_Name}")]
    if start <= story.created <= end:
        lines.append((story.created, "Issue created"))
    window = events[(events["Creation_Date"] >= start) & (events["Creation_Date"] <= end)]
    for row in window.itertuples():
        line = describe(row.Field, row.From_String, row.To_String)
        if line:
            lines.append((row.Creation_Date, line))
    lines.sort(key=lambda item: item[0])
    text = [f"{when.strftime(fmt)}  {line}" for when, line in lines[:limit]]
    if len(lines) > limit:
        text.append(f"… {len(lines) - limit} more events (see Jira)")
    return "\n".join(text)


def review_rows(sample: pd.DataFrame, snapshot: pd.DataFrame) -> pd.DataFrame:
    ids = sample["Issue_ID"].tolist()
    stories = sample[["Issue_ID", "stratum"]].merge(snapshot, on="Issue_ID", how="left")
    urls = tawos.load("Issue", ["ID", "URL"], filters=pc.field("ID").isin(ids)).set_index("ID")["URL"]
    events = tawos.load("Change_Log", ["ID", "Issue_ID", "Field", "From_String", "To_String", "Creation_Date"],
                       filters=pc.field("Field").isin(["Sprint", "status", "resolution", "Story Points", "Link"])
                       & pc.field("Issue_ID").isin(ids)).sort_values(["Creation_Date", "ID"])
    by_issue = dict(tuple(events.groupby("Issue_ID")))
    empty = events.iloc[0:0]
    fmt = "%Y-%m-%d %H:%M"
    return pd.DataFrame({
        "Issue_ID": stories["Issue_ID"],
        "Issue": stories["Issue_Key"],
        "url": [browse_url(urls.get(i), k) for i, k in zip(stories["Issue_ID"], stories["Issue_Key"], strict=True)],
        "Project": stories["Project_Key"],
        "Type": stories["issue_type"],
        "Points when committed": stories["story_points"],
        "Title": stories["title"].str.slice(0, 200),
        "First sprint": stories["Sprint_Name"],
        "Sprint started": stories["sprint_start"].dt.strftime(fmt),
        "Sprint closed": stories["sprint_closed"].dt.strftime(fmt),
        "How it entered": ["Planned (in the sprint at its start)" if not mid
                           else f"Added during the sprint, {when.strftime(fmt)}"
                           for mid, when in zip(stories["added_mid_sprint"], stories["commitment_time"], strict=True)],
        "What happened": [history_text(by_issue.get(s.Issue_ID, empty), s) for s in stories.itertuples()],
    })


def write_workbook(rows: pd.DataFrame, path: Path, reviewer: str) -> None:
    wb = Workbook()
    guide = wb.active
    guide.title = "How to review"
    stories = wb.create_sheet("Stories")
    bold, plain = Font(name=FONT, bold=True), Font(name=FONT)
    fill = PatternFill("solid", start_color="FFFF99")
    wrap = Alignment(wrap_text=True, vertical="top")

    lines = [
        (f"Synapse effort and risk predictor: checking the automatic risk labels (reviewer {reviewer})", True),
        ("", False),
        ("Why: the risk model learns from labels computed automatically from Jira history. Before it is trained, "
         "two people check 200 stories by hand, each on their own. Their agreement is reported in the thesis.", False),
        ("", False),
        ("What to do for each row of the Stories sheet:", True),
        ("1. Read 'What happened' (and open the Jira link if you want more detail; some old Jira sites are offline).",
         False),
        ("2. Decide whether the story really ran into trouble in its first sprint, using the six signs below.", False),
        ("3. Fill in ONLY the yellow columns: At risk? (Yes/No), Signs you saw, Confidence (1-3), Notes.", False),
        ("Work on your own and do not compare answers with the other reviewer until both of you have finished.", True),
        ("", False),
        ("The six warning signs (a story is at risk if at least one really happened):", True),
        *[(f"{code} {name}: {meaning}", False) for code, name, meaning in SIGNS],
        ("", False),
        ("Use your judgement where the history looks odd, for example a story moved only because the sprint was "
         "reorganised, or points changed only to fix a typo. Say why in Notes.", False),
        ("Times: sprint start and close times can be a few hours off from the other times (different clocks in "
         "the source data), so do not decide on a difference of a few hours alone.", False),
        ("", False),
        ("Example of a filled-in row (not one of the 200 stories):", True),
    ]
    for i, (text, is_bold) in enumerate(lines, start=1):
        cell = guide.cell(row=i, column=1, value=text)
        cell.font, cell.alignment = (bold if is_bold else plain), Alignment(wrap_text=True, vertical="top")
    example_headers = ["Issue", "At risk? (Yes / No)", "Signs you saw (R1–R6)", "Confidence (1–3)", "Notes"]
    example = ["ABC-123", "Yes", "R1, R5", 3, "Not resolved at the close; moved to the next sprint the same day."]
    top = len(lines) + 1
    for col, (header, value) in enumerate(zip(example_headers, example, strict=True), start=1):
        guide.cell(row=top, column=col, value=header).font = bold
        cell = guide.cell(row=top + 1, column=col, value=value)
        cell.font = plain
        if col > 1:
            cell.fill = fill
    guide.cell(row=top + 3, column=1, value="Reviewer name:").font = bold
    guide.cell(row=top + 3, column=2).fill = fill
    guide.column_dimensions["A"].width = 110
    for letter in "BCDE":
        guide.column_dimensions[letter].width = 22

    info = ["#", "Issue", "Project", "Type", "Points when committed", "Title", "First sprint", "Sprint started",
            "Sprint closed", "How it entered", "What happened (3 days before the sprint to the end of the next one)"]
    inputs = ["At risk? (Yes / No)", "Signs you saw (R1–R6)", "Confidence (1 = unsure, 3 = sure)", "Notes"]
    widths = [5, 15, 12, 14, 13, 40, 26, 17, 17, 24, 90, 12, 16, 14, 36]
    thin = Side(style="thin", color="BBBBBB")
    for col, (header, width) in enumerate(zip(info + inputs, widths, strict=True), start=1):
        cell = stories.cell(row=1, column=col, value=header)
        cell.font, cell.alignment = bold, Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=thin)
        if header in inputs:
            cell.fill = fill
        stories.column_dimensions[cell.column_letter].width = width
    stories.row_dimensions[1].height = 45

    columns = ["Issue", "Project", "Type", "Points when committed", "Title", "First sprint", "Sprint started",
               "Sprint closed", "How it entered", "What happened"]
    for r, row in enumerate(rows.to_dict("records"), start=2):
        for col, value in enumerate([r - 1, *[row[c] for c in columns]], start=1):
            cell = stories.cell(row=r, column=col, value=value)
            cell.font, cell.alignment = plain, wrap
        if row["url"]:
            cell = stories.cell(row=r, column=2)
            cell.hyperlink, cell.font = row["url"], Font(name=FONT, color="0563C1", underline="single")
        for col in range(len(info) + 1, len(info) + len(inputs) + 1):
            cell = stories.cell(row=r, column=col)
            cell.fill, cell.font, cell.alignment = fill, plain, wrap
        lines = max(row["What happened"].count("\n") + 1, len(str(row["Title"])) // 38 + 1)
        stories.row_dimensions[r].height = min(409, 15 * lines + 6)  # points; about 15 per line of Arial 11

    last = len(rows) + 1
    yes_no = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    yes_no.add(f"L2:L{last}")
    confidence = DataValidation(type="whole", operator="between", formula1="1", formula2="3", allow_blank=True)
    confidence.add(f"N2:N{last}")
    stories.add_data_validation(yes_no)
    stories.add_data_validation(confidence)
    stories.freeze_panes = "C2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--out-dir", type=Path, default=config.WORK_DIR / "review")
    parser.add_argument("--size", type=int, default=200)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--force", action="store_true", help="overwrite workbooks that were already handed out")
    args = parser.parse_args(argv)
    existing = sorted(args.out_dir.glob("label-review-*.xlsx"))
    if existing and not args.force:
        raise SystemExit(f"{existing[0].parent} already holds review workbooks, maybe with answers in them. "
                         "Use --force to draw a new sample and overwrite them.")

    labels = pd.read_parquet(config.INTERIM_DIR / "labels.parquet")
    snapshot = pd.read_parquet(config.INTERIM_DIR / "snapshot.parquet")
    sample = sample_for_review(labels, n=args.size, seed=args.seed)
    rows = review_rows(sample, snapshot)
    for reviewer in REVIEWERS:
        write_workbook(rows, args.out_dir / f"label-review-reviewer-{reviewer}.xlsx", reviewer)
    key = sample.assign(row=range(1, len(sample) + 1))
    key.to_parquet(args.out_dir / "label-review-key.parquet", index=False)
    print(f"Wrote {len(rows)} stories for reviewers {', '.join(REVIEWERS)} and the answer key to {args.out_dir}")


if __name__ == "__main__":
    main()
